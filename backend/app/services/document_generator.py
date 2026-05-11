import uuid
from datetime import date
from pathlib import Path

from docx import Document


class DocumentGenerator:
    def __init__(self, documents_dir: Path, temp_dir: Path):
        self.documents_dir = Path(documents_dir)
        self.temp_dir = Path(temp_dir)
        self.temp_dir.mkdir(parents=True, exist_ok=True)

    def generate(self, doc_type: str, fields: dict[str, str], template_name: str = "template.docx") -> Path:
        """Fill template with field values. template_name can be overridden per variant (e.g. 'template_course_teacher.docx')."""
        template_path = self.documents_dir / doc_type / template_name
        if not template_path.exists():
            raise FileNotFoundError(f"Template not found: {template_path}")

        doc = Document(str(template_path))

        # Auto-fill sign_date if not provided
        if "sign_date" not in fields or not fields.get("sign_date"):
            fields = {**fields, "sign_date": date.today().strftime("%Y年%m月%d日")}

        # Replace placeholders in paragraphs
        for paragraph in doc.paragraphs:
            self._replace_in_paragraph(paragraph, fields)

        # Replace placeholders in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._replace_in_paragraph(paragraph, fields)

        output_path = self.temp_dir / f"{doc_type}_{uuid.uuid4().hex[:12]}.docx"
        doc.save(str(output_path))
        return output_path

    @staticmethod
    def _replace_in_paragraph(paragraph, fields: dict[str, str]) -> None:
        for key, value in fields.items():
            placeholder = f"{{{{{key}}}}}"
            display_value = value if value else "________"
            for run in paragraph.runs:
                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, display_value)
