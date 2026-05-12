import io
from pathlib import Path
import pytest
from docx import Document


@pytest.fixture
def templates_dir(tmp_path):
    doc_dir = tmp_path / "leave"
    doc_dir.mkdir(parents=True)
    doc = Document()
    doc.add_paragraph("Name: {{name}}")
    doc.add_paragraph("ID: {{student_id}}")
    doc.save(str(doc_dir / "template.docx"))
    return tmp_path


def test_fills_placeholders(templates_dir):
    from app.services.document_generator import DocumentGenerator

    gen = DocumentGenerator(templates_dir, templates_dir / "tmp")
    output = gen.generate("leave", {"name": "张三", "student_id": "20230001"})

    assert output.exists()
    result = Document(str(output))
    text = "\n".join(p.text for p in result.paragraphs)
    assert "张三" in text
    assert "20230001" in text
    assert "{{name}}" not in text


def test_unfilled_placeholders_become_blanks(templates_dir):
    from app.services.document_generator import DocumentGenerator

    gen = DocumentGenerator(templates_dir, templates_dir / "tmp")
    output = gen.generate("leave", {})

    result = Document(str(output))
    text = "\n".join(p.text for p in result.paragraphs)
    assert "________" in text


def test_sign_date_auto_filled(templates_dir):
    from datetime import date
    from app.services.document_generator import DocumentGenerator

    doc_dir = templates_dir / "leave2"
    doc_dir.mkdir(parents=True)
    doc = Document()
    doc.add_paragraph("Date: {{sign_date}}")
    doc.save(str(doc_dir / "template.docx"))

    gen = DocumentGenerator(templates_dir, templates_dir / "tmp")
    output = gen.generate("leave2", {"reason": "sick"})

    result = Document(str(output))
    text = "\n".join(p.text for p in result.paragraphs)
    today_str = date.today().strftime("%Y年%m月%d日")
    assert today_str in text
